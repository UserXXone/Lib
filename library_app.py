import csv
import shutil
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from pathlib import Path
from datetime import datetime, timedelta
import time
import re
import zipfile

DATA_PATH = Path("kutuphane.csv")
BACKUP_DIR = Path("yedekler")
STUDENT_PATH = Path("ogrenciler.csv")
COLUMNS = ["KitapID", "QR", "KitapAdi", "Yazar", "Raf", "Durum", "OduncAlan", "TeslimTarihi"]
STUDENT_COLUMNS = ["OgrenciNo", "AdSoyad"]
GUNLER = ["Pazartesi", "Salı", "Çarşamba", "Perşembe", "Cuma", "Cumartesi", "Pazar"]

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


class LibraryData:
    def __init__(self, file_path: Path):
        self.file_path = file_path
        self._ensure_file()

    def _ensure_file(self):
        if self.file_path.exists():
            return
        rows = [
            {"KitapID": "1", "QR": "LIB-000001", "KitapAdi": "Suç ve Ceza", "Yazar": "Dostoyevski", "Raf": "A-1", "Durum": "Rafta", "OduncAlan": "", "TeslimTarihi": ""},
            {"KitapID": "2", "QR": "LIB-000002", "KitapAdi": "Kürk Mantolu Madonna", "Yazar": "Sabahattin Ali", "Raf": "A-1", "Durum": "Ödünçte", "OduncAlan": "Ayşe", "TeslimTarihi": "2026-01-15"},
            {"KitapID": "3", "QR": "LIB-000003", "KitapAdi": "1984", "Yazar": "George Orwell", "Raf": "B-2", "Durum": "Rafta", "OduncAlan": "", "TeslimTarihi": ""},
            {"KitapID": "4", "QR": "LIB-000004", "KitapAdi": "Simyacı", "Yazar": "Paulo Coelho", "Raf": "C-3", "Durum": "Rafta", "OduncAlan": "", "TeslimTarihi": ""},
        ]
        self.write(rows)

    def read(self):
        rows = []
        with self.file_path.open("r", newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            for row in reader:
                rows.append({col: row.get(col, "") for col in COLUMNS})
        return self._ensure_qr_values(rows)

    def write(self, rows):
        with self.file_path.open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=COLUMNS)
            writer.writeheader()
            writer.writerows(rows)

    def _make_qr(self, book_id: int):
        return f"LIB-{book_id:06d}"

    def _normalize_import_qr(self, qr: str):
        value = str(qr).strip()
        if not value:
            return ""
        # Import isteğine göre QR içindeki tüm tireleri yıldız yap
        return value.replace("-", "*")

    def _identifier_matches(self, row, code: str):
        code = str(code).strip()
        return code == str(row.get("KitapID", "")).strip() or code.lower() == str(row.get("QR", "")).strip().lower()

    def _ensure_qr_values(self, rows):
        changed = False
        for row in rows:
            if not str(row.get("QR", "")).strip() and str(row.get("KitapID", "")).strip().isdigit():
                row["QR"] = self._make_qr(int(row["KitapID"]))
                changed = True
        if changed:
            self.write(rows)
        return rows

    def _find_row(self, rows, code: str):
        return next((r for r in rows if self._identifier_matches(r, code)), None)

    def add_book(self, title: str, author: str, shelf: str):
        rows = self.read()
        max_id = max((int(r["KitapID"]) for r in rows), default=0)
        rows.append(
            {
                "KitapID": str(max_id + 1),
                "QR": self._make_qr(max_id + 1),
                "KitapAdi": title.strip(),
                "Yazar": author.strip(),
                "Raf": shelf.strip(),
                "Durum": "Rafta",
                "OduncAlan": "",
                "TeslimTarihi": "",
            }
        )
        self.write(rows)

    def update_book(self, code: str, title: str, author: str, shelf: str):
        rows = self.read()
        row = self._find_row(rows, code)
        if not row:
            raise ValueError("Kitap bulunamadı (QR/ID).")
        row["KitapAdi"] = title.strip()
        row["Yazar"] = author.strip()
        row["Raf"] = shelf.strip()
        self.write(rows)

    def delete_book(self, code: str):
        rows = self.read()
        row = self._find_row(rows, code)
        if not row:
            raise ValueError("Kitap bulunamadı (QR/ID).")
        if row["Durum"] == "Ödünçte":
            raise ValueError("Ödünçte olan kitap silinemez, önce iade alın.")
        self.write([r for r in rows if not self._identifier_matches(r, code)])

    def lend_book(self, code: str, borrower: str, days: int = 14):
        rows = self.read()
        row = self._find_row(rows, code)
        if not row:
            raise ValueError("Kitap bulunamadı (QR/ID).")
        if row["Durum"] == "Ödünçte":
            raise ValueError("Bu kitap zaten ödünçte.")
        row["Durum"] = "Ödünçte"
        row["OduncAlan"] = borrower.strip()
        row["TeslimTarihi"] = (datetime.now() + timedelta(days=days)).strftime("%Y-%m-%d")
        self.write(rows)

    def return_book(self, code: str):
        rows = self.read()
        row = self._find_row(rows, code)
        if not row:
            raise ValueError("Kitap bulunamadı (QR/ID).")
        if row["Durum"] != "Ödünçte":
            raise ValueError("Bu kitap zaten rafta görünüyor.")
        row["Durum"] = "Rafta"
        row["OduncAlan"] = ""
        row["TeslimTarihi"] = ""
        self.write(rows)

    def extend_due_date(self, code: str, days: int):
        rows = self.read()
        row = self._find_row(rows, code)
        if not row:
            raise ValueError("Kitap bulunamadı (QR/ID).")
        if row["Durum"] != "Ödünçte":
            raise ValueError("Sadece ödünçte olan kitapların süresi uzatılabilir.")
        base = datetime.now().date()
        if row["TeslimTarihi"]:
            try:
                base = datetime.strptime(row["TeslimTarihi"], "%Y-%m-%d").date()
            except ValueError:
                pass
        row["TeslimTarihi"] = (base + timedelta(days=days)).strftime("%Y-%m-%d")
        self.write(rows)

    def overdue_books(self):
        today = datetime.now().date()
        overdues = []
        for row in self.read():
            if row["Durum"] != "Ödünçte" or not row["TeslimTarihi"]:
                continue
            try:
                due = datetime.strptime(row["TeslimTarihi"], "%Y-%m-%d").date()
            except ValueError:
                continue
            if due < today:
                overdues.append((row, (today - due).days))
        return overdues

    def due_today(self):
        today = datetime.now().strftime("%Y-%m-%d")
        return [r for r in self.read() if r["Durum"] == "Ödünçte" and r["TeslimTarihi"] == today]

    def stats(self):
        rows = self.read()
        return {
            "Toplam": len(rows),
            "Rafta": len([r for r in rows if r["Durum"] == "Rafta"]),
            "Ödünçte": len([r for r in rows if r["Durum"] == "Ödünçte"]),
            "Geciken": len(self.overdue_books()),
        }


    def _read_import_source(self, file_path: Path):
        ext = file_path.suffix.lower()

        if ext == ".csv":
            with file_path.open("r", newline="", encoding="utf-8-sig") as f:
                reader = csv.DictReader(f)
                headers = list(reader.fieldnames or [])
                rows = [dict(r) for r in reader]
            return headers, rows

        if ext in (".xlsx", ".xlsm"):
            if load_workbook is None:
                raise ValueError("XLSX import için openpyxl gerekli. CSV içe aktarımı kullanabilirsiniz.")
            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            it = ws.iter_rows(values_only=True)
            header_row = next(it, None)
            if not header_row:
                return [], []
            headers = [str(h).strip() if h is not None else "" for h in header_row]
            rows = []
            for r in it:
                row = {}
                for i, h in enumerate(headers):
                    if not h:
                        continue
                    val = r[i] if i < len(r) else ""
                    row[h] = "" if val is None else str(val)
                rows.append(row)
            return headers, rows

        raise ValueError("Desteklenmeyen dosya türü. CSV veya XLSX seçin.")

    def get_import_headers(self, file_path: Path):
        headers, _ = self._read_import_source(file_path)
        return [h for h in headers if h]

    def suggest_import_mapping(self, headers):
        norm = {h.strip().lower(): h for h in headers}

        def pick(candidates):
            for c in candidates:
                if c in norm:
                    return norm[c]
            return ""

        return {
            "title": pick(["kitapadi", "kitap adı", "aciklama", "açıklama", "orijinal_adi", "orijinal adı"]),
            "author": pick(["yazar", "hazirlayan", "hazırlayan", "cevirmen", "çevirmen", "duzeltilmis kategori", "duzeltilmis_kategori"]),
            "shelf": pick(["raf", "RAF", "raf_no"]),
            "qr": pick(["qr", "QR", "barkod"]),
        }

    def import_books(self, file_path: Path, mapping: dict | None = None):
        headers, rows = self._read_import_source(file_path)
        if not headers:
            return 0

        mapping = mapping or self.suggest_import_mapping(headers)
        title_col = mapping.get("title", "")
        author_col = mapping.get("author", "")
        shelf_col = mapping.get("shelf", "")
        qr_col = mapping.get("qr", "")

        if not title_col or not shelf_col:
            raise ValueError("Import eşleştirmesinde en az Kitap Adı ve Raf sütunları seçilmelidir.")

        imported = 0
        current_rows = self.read()
        max_id = max((int(r["KitapID"]) for r in current_rows), default=0)

        for row in rows:
            title = str(row.get(title_col, "")).strip()
            author = str(row.get(author_col, "")).strip() if author_col else "Bilinmiyor"
            shelf = str(row.get(shelf_col, "")).strip()
            qr_val = self._normalize_import_qr(row.get(qr_col, "")) if qr_col else ""
            if title and shelf:
                max_id += 1
                current_rows.append(
                    {
                        "KitapID": str(max_id),
                        "QR": qr_val or self._make_qr(max_id),
                        "KitapAdi": title,
                        "Yazar": author or "Bilinmiyor",
                        "Raf": shelf,
                        "Durum": "Rafta",
                        "OduncAlan": "",
                        "TeslimTarihi": "",
                    }
                )
                imported += 1

        self.write(current_rows)
        return imported

    def clear_database(self):
        self.write([])

    def create_backup(self):
        BACKUP_DIR.mkdir(exist_ok=True)
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        package_dir = BACKUP_DIR / f"tam_yedek_{ts}"
        package_dir.mkdir(parents=True, exist_ok=True)

        # Uygulamadaki temel her şeyi paketle
        include_files = [Path("library_app.py"), Path("README.md"), Path("requirements.txt"), DATA_PATH, STUDENT_PATH]
        for file_path in include_files:
            if file_path.exists():
                shutil.copy2(file_path, package_dir / file_path.name)

        # Kök dizindeki diğer CSV'leri de ekle
        for csv_file in Path('.').glob('*.csv'):
            target = package_dir / csv_file.name
            if csv_file.exists() and not target.exists():
                shutil.copy2(csv_file, target)

        zip_target = BACKUP_DIR / f"tam_yedek_{ts}.zip"
        with zipfile.ZipFile(zip_target, 'w', zipfile.ZIP_DEFLATED) as zf:
            for fp in package_dir.glob('*'):
                zf.write(fp, arcname=fp.name)

        shutil.rmtree(package_dir, ignore_errors=True)
        return zip_target


class StudentData:
    def __init__(self, file_path: Path):
        self.file_path = file_path
        self._ensure_file()

    def _ensure_file(self):
        if self.file_path.exists():
            return
        self.write([])

    def read(self):
        rows = []
        with self.file_path.open('r', newline='', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                rows.append({col: row.get(col, '') for col in STUDENT_COLUMNS})
        return rows

    def write(self, rows):
        with self.file_path.open('w', newline='', encoding='utf-8') as f:
            writer = csv.DictWriter(f, fieldnames=STUDENT_COLUMNS)
            writer.writeheader()
            writer.writerows(rows)

    def add_student(self, no: str, name: str):
        no = str(no).strip()
        if not no:
            raise ValueError('Öğrenci numarası boş olamaz.')
        rows = self.read()
        if any(str(r['OgrenciNo']) == no for r in rows):
            raise ValueError('Bu öğrenci numarası zaten kayıtlı.')
        rows.append({'OgrenciNo': no, 'AdSoyad': name.strip()})
        self.write(rows)
        return no

    def update_student(self, no: str, name: str):
        rows = self.read()
        for row in rows:
            if str(row['OgrenciNo']) == str(no):
                row['AdSoyad'] = name.strip()
                self.write(rows)
                return
        raise ValueError('Öğrenci numarası bulunamadı.')

    def delete_student(self, no: str):
        rows = self.read()
        if not any(str(r['OgrenciNo']) == str(no) for r in rows):
            raise ValueError('Öğrenci numarası bulunamadı.')
        rows = [r for r in rows if str(r['OgrenciNo']) != str(no)]
        self.write(rows)


class LibraryApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Kütüphane Takip Sistemi")
        self.geometry("1220x720")
        self.minsize(1000, 620)
        self.configure(bg="#f4f6fb")

        self.data = LibraryData(DATA_PATH)
        self.students = StudentData(STUDENT_PATH)
        self.clock_var = tk.StringVar()

        self._setup_styles()
        self._build_layout()
        self._start_clock()
        self.show_home()

        # Kısayol: Alt+1 ile veritabanını temizle
        self.bind_all("<Alt-Key-1>", self._handle_db_clear_shortcut)

    def _setup_styles(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("Card.TFrame", background="white")
        style.configure("Title.TLabel", font=("Segoe UI", 17, "bold"), background="#f4f6fb", foreground="#111827")
        style.configure("SubTitle.TLabel", font=("Segoe UI", 11), background="#f4f6fb", foreground="#6b7280")
        style.configure("CardTitle.TLabel", font=("Segoe UI", 10), background="white", foreground="#64748b")
        style.configure("CardValue.TLabel", font=("Segoe UI", 18, "bold"), background="white", foreground="#0f172a")

    def _build_layout(self):
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=1)

        self.menu_frame = tk.Frame(self, bg="#0f172a", width=250)
        self.menu_frame.grid(row=0, column=0, sticky="ns")
        self.menu_frame.grid_propagate(False)

        tk.Label(
            self.menu_frame,
            text="Kütüphane\nPaneli",
            fg="white",
            bg="#0f172a",
            font=("Segoe UI", 18, "bold"),
            justify="center",
        ).pack(pady=(24, 20))

        self._menu_button("Ana Ekran", self.show_home)
        self._menu_button("Kitap Sorgula", self.show_query)
        self._menu_button("Kitap Ekle", self.show_add_book)
        self._menu_button("Düzenle / Sil", self.show_manage_book)
        self._menu_button("Öğrenci Paneli", self.show_students_panel)
        self._menu_button("Ödünç Ver", self.show_lend)
        self._menu_button("İade Al", self.show_return)
        self._menu_button("Süre Uzat", self.show_extend_due)
        self._menu_button("Gecikenler", self.show_overdues)
        self._menu_button("Yedek Al", self.make_backup)
        self._menu_button("Çıkış", self.exit_app)

        self.main_frame = tk.Frame(self, bg="#f4f6fb")
        self.main_frame.grid(row=0, column=1, sticky="nsew")
        self.main_frame.grid_columnconfigure(0, weight=1)
        self.main_frame.grid_rowconfigure(0, weight=1)

        self.status_bar = tk.Frame(self.main_frame, bg="#e2e8f0", height=34)
        self.status_bar.grid(row=1, column=0, sticky="ew")
        self.status_bar.grid_columnconfigure(0, weight=1)

        tk.Label(
            self.status_bar,
            text="Sürüm 1.1 • Modern görünüm etkin",
            bg="#e2e8f0",
            fg="#374151",
            font=("Segoe UI", 9),
        ).grid(row=0, column=0, sticky="w", padx=12)

        tk.Label(
            self.status_bar,
            textvariable=self.clock_var,
            bg="#e2e8f0",
            fg="#111827",
            font=("Segoe UI", 10, "bold"),
        ).grid(row=0, column=1, sticky="e", padx=12)

    def _menu_button(self, text, command):
        tk.Button(
            self.menu_frame,
            text=text,
            command=command,
            bg="#1e293b",
            fg="white",
            activebackground="#334155",
            activeforeground="white",
            bd=0,
            padx=16,
            pady=9,
            font=("Segoe UI", 11),
            anchor="w",
            cursor="hand2",
        ).pack(fill="x", padx=12, pady=4)

    def _start_clock(self):
        now = datetime.now()
        text = now.strftime("%d.%m.%Y") + f" {GUNLER[now.weekday()]} • " + now.strftime("%H:%M:%S")
        self.clock_var.set(text)
        self.after(1000, self._start_clock)

    def _clear_main(self):
        for widget in self.main_frame.grid_slaves(row=0, column=0):
            widget.destroy()

    def _stats_cards(self, parent):
        stats = self.data.stats()
        row = ttk.Frame(parent)
        row.grid(row=0, column=0, sticky="ew", pady=(0, 12))
        for i in range(4):
            row.grid_columnconfigure(i, weight=1)

        icons = {"Toplam": "•", "Rafta": "•", "Ödünçte": "•", "Geciken": "•"}
        for i, key in enumerate(["Toplam", "Rafta", "Ödünçte", "Geciken"]):
            card = ttk.Frame(row, style="Card.TFrame", padding=12)
            card.grid(row=0, column=i, sticky="ew", padx=6)
            ttk.Label(card, text=f"{icons[key]} {key}", style="CardTitle.TLabel").pack(anchor="w")
            ttk.Label(card, text=str(stats[key]), style="CardValue.TLabel").pack(anchor="w")

    def show_home(self):
        self._clear_main()
        wrapper = ttk.Frame(self.main_frame, padding=20)
        wrapper.grid(row=0, column=0, sticky="nsew")
        wrapper.grid_columnconfigure(0, weight=3)
        wrapper.grid_columnconfigure(1, weight=2)
        wrapper.grid_rowconfigure(2, weight=1)

        self._stats_cards(wrapper)
        ttk.Label(wrapper, text="Raf Durumu", style="Title.TLabel").grid(row=1, column=0, sticky="w", pady=(0, 8))

        rows = self.data.read()
        raf_ozet = {}
        for r in rows:
            raf = r["Raf"] or "Belirsiz"
            raf_ozet.setdefault(raf, {"Rafta": 0, "Ödünçte": 0})
            durum = r["Durum"] if r["Durum"] in ("Rafta", "Ödünçte") else "Rafta"
            raf_ozet[raf][durum] += 1

        tree = ttk.Treeview(wrapper, columns=("Raf", "Rafta", "Ödünçte", "Toplam"), show="headings", height=15)
        for col in ("Raf", "Rafta", "Ödünçte", "Toplam"):
            tree.heading(col, text=col)
            tree.column(col, width=120, anchor="center")
        for raf, counts in sorted(raf_ozet.items()):
            total = counts["Rafta"] + counts["Ödünçte"]
            tree.insert("", "end", values=(raf, counts["Rafta"], counts["Ödünçte"], total))
        tree.grid(row=2, column=0, sticky="nsew", padx=(0, 10))

        side = ttk.Frame(wrapper, style="Card.TFrame", padding=14)
        side.grid(row=2, column=1, sticky="nsew")
        ttk.Label(side, text="Bugün Teslim Edilecekler", style="CardTitle.TLabel").pack(anchor="w")

        due_today = self.data.due_today()
        if not due_today:
            ttk.Label(side, text="Bugün teslim beklenen kitap yok.", background="white", foreground="#374151").pack(anchor="w", pady=(6, 8))
        else:
            for item in due_today[:8]:
                ttk.Label(side, text=f"• {item['KitapAdi']} ({item['OduncAlan']})", background="white").pack(anchor="w", pady=2)

        ttk.Label(side, text="\nİpucu: Gecikenleri menüden hızlıca görüntüleyebilirsin.", background="white", foreground="#64748b").pack(anchor="w")

    def _standard_form_title(self, wrapper, title, subtitle=""):
        ttk.Label(wrapper, text=title, style="Title.TLabel").pack(anchor="w")
        if subtitle:
            ttk.Label(wrapper, text=subtitle, style="SubTitle.TLabel").pack(anchor="w", pady=(2, 10))

    def _run_with_progress(self, title, action):
        dlg = tk.Toplevel(self)
        dlg.title(title)
        dlg.transient(self)
        dlg.grab_set()
        dlg.resizable(False, False)

        ttk.Label(dlg, text="İşlem yapılıyor...", font=("Segoe UI", 10, "bold")).pack(anchor="w", padx=12, pady=(10, 6))
        percent_var = tk.StringVar(value="0%")
        bar = ttk.Progressbar(dlg, orient="horizontal", length=340, mode="determinate", maximum=100)
        bar.pack(padx=12, pady=4)
        ttk.Label(dlg, textvariable=percent_var).pack(anchor="e", padx=12, pady=(0, 10))

        def step(value):
            bar["value"] = value
            percent_var.set(f"{int(value)}%")
            dlg.update_idletasks()
            time.sleep(0.08)

        result = {"ok": False, "err": None}
        try:
            for pct in (15, 45, 75):
                step(pct)
            action()
            step(100)
            result["ok"] = True
        except Exception as err:
            result["err"] = err
        finally:
            dlg.destroy()

        if result["err"]:
            raise result["err"]

    def show_query(self):
        self._clear_main()
        wrapper = ttk.Frame(self.main_frame, padding=20)
        wrapper.grid(row=0, column=0, sticky="nsew")
        self._standard_form_title(wrapper, "Kitap Sorgulama", "Metin araması, durum filtresi ve hızlı yenileme")

        search_var = tk.StringVar()
        durum_var = tk.StringVar(value="Hepsi")

        top = ttk.Frame(wrapper)
        top.pack(fill="x", pady=8)
        ttk.Entry(top, textvariable=search_var, width=40).pack(side="left", padx=(0, 8))
        ttk.Combobox(top, textvariable=durum_var, state="readonly", values=["Hepsi", "Rafta", "Ödünçte"], width=14).pack(side="left")

        tree = ttk.Treeview(wrapper, columns=COLUMNS, show="headings", height=18)
        for col in COLUMNS:
            tree.heading(col, text=col)
            tree.column(col, width=124, anchor="center")
        tree.pack(fill="both", expand=True, pady=(10, 0))

        def load_rows(*_):
            tree.delete(*tree.get_children())
            q = search_var.get().strip().lower()
            durum_filter = durum_var.get()
            for row in self.data.read():
                if durum_filter != "Hepsi" and row["Durum"] != durum_filter:
                    continue
                haystack = f"{row['KitapAdi']} {row['Yazar']} {row['Raf']} {row['OduncAlan']} {row.get('QR','')} {row.get('KitapID','')}".lower()
                if q and q not in haystack:
                    continue
                tree.insert("", "end", values=[row[c] for c in COLUMNS])

        ttk.Button(top, text="Yenile", command=load_rows).pack(side="left", padx=8)
        search_var.trace_add("write", load_rows)
        durum_var.trace_add("write", load_rows)
        load_rows()

    def _open_import_mapping_dialog(self, file_path: Path):
        headers = self.data.get_import_headers(file_path)
        if not headers:
            raise ValueError("Dosyada başlık satırı bulunamadı.")

        suggested = self.data.suggest_import_mapping(headers)
        choices = [""] + headers

        dlg = tk.Toplevel(self)
        dlg.title("Import Sütun Eşleştirme")
        dlg.transient(self)
        dlg.grab_set()
        dlg.resizable(False, False)

        ttk.Label(dlg, text="Dosyadaki sütunları alanlara eşleştir:", font=("Segoe UI", 10, "bold")).grid(row=0, column=0, columnspan=2, padx=12, pady=(10, 8), sticky="w")

        title_var = tk.StringVar(value=suggested.get("title", ""))
        author_var = tk.StringVar(value=suggested.get("author", ""))
        shelf_var = tk.StringVar(value=suggested.get("shelf", ""))
        qr_var = tk.StringVar(value=suggested.get("qr", ""))

        ttk.Label(dlg, text="Kitap Adı sütunu:").grid(row=1, column=0, padx=12, pady=5, sticky="w")
        ttk.Combobox(dlg, textvariable=title_var, values=choices, state="readonly", width=36).grid(row=1, column=1, padx=12, pady=5)

        ttk.Label(dlg, text="Yazar sütunu:").grid(row=2, column=0, padx=12, pady=5, sticky="w")
        ttk.Combobox(dlg, textvariable=author_var, values=choices, state="readonly", width=36).grid(row=2, column=1, padx=12, pady=5)

        ttk.Label(dlg, text="Raf sütunu:").grid(row=3, column=0, padx=12, pady=5, sticky="w")
        ttk.Combobox(dlg, textvariable=shelf_var, values=choices, state="readonly", width=36).grid(row=3, column=1, padx=12, pady=5)

        ttk.Label(dlg, text="QR sütunu (opsiyonel):").grid(row=4, column=0, padx=12, pady=5, sticky="w")
        ttk.Combobox(dlg, textvariable=qr_var, values=choices, state="readonly", width=36).grid(row=4, column=1, padx=12, pady=5)

        result = {"confirmed": False, "mapping": None}

        def use_sample_template():
            # Gönderdiğiniz şablona göre öneri
            if "aciklama" in headers:
                title_var.set("aciklama")
            if "DUZELTILMIS KATEGORI" in headers:
                author_var.set("DUZELTILMIS KATEGORI")
            elif "hazirlayan" in headers:
                author_var.set("hazirlayan")
            if "RAF" in headers:
                shelf_var.set("RAF")
            elif "raf" in headers:
                shelf_var.set("raf")
            if "QR" in headers:
                qr_var.set("QR")
            elif "qr" in headers:
                qr_var.set("qr")

        def confirm():
            if not title_var.get().strip() or not shelf_var.get().strip():
                messagebox.showwarning("Eksik eşleştirme", "Kitap Adı ve Raf sütunu zorunludur.", parent=dlg)
                return
            result["confirmed"] = True
            result["mapping"] = {
                "title": title_var.get().strip(),
                "author": author_var.get().strip(),
                "shelf": shelf_var.get().strip(),
                "qr": qr_var.get().strip(),
            }
            dlg.destroy()

        btns = ttk.Frame(dlg)
        btns.grid(row=5, column=0, columnspan=2, padx=12, pady=12, sticky="e")
        ttk.Button(btns, text="Örnek Şablonu Uygula", command=use_sample_template).grid(row=0, column=0, padx=(0, 8))
        ttk.Button(btns, text="İptal", command=dlg.destroy).grid(row=0, column=1, padx=(0, 8))
        ttk.Button(btns, text="Import Et", command=confirm).grid(row=0, column=2)

        self.wait_window(dlg)
        return result["mapping"] if result["confirmed"] else None

    def show_add_book(self):
        self._clear_main()
        wrapper = ttk.Frame(self.main_frame, padding=20)
        wrapper.grid(row=0, column=0, sticky="nsew")
        self._standard_form_title(wrapper, "Yeni Kitap Ekle", "Kayıt sonrası kitap otomatik olarak rafta görünür")

        title_var, author_var, shelf_var = tk.StringVar(), tk.StringVar(), tk.StringVar()
        form = ttk.Frame(wrapper)
        form.pack(anchor="w", pady=10)

        for i, (label, var) in enumerate([("Kitap Adı", title_var), ("Yazar", author_var), ("Raf", shelf_var)]):
            ttk.Label(form, text=f"{label}:").grid(row=i, column=0, sticky="w", pady=6)
            ttk.Entry(form, textvariable=var, width=42).grid(row=i, column=1, padx=8, pady=6)

        def save():
            if not title_var.get().strip() or not author_var.get().strip() or not shelf_var.get().strip():
                messagebox.showwarning("Eksik bilgi", "Lütfen tüm alanları doldurun.")
                return
            self.data.add_book(title_var.get(), author_var.get(), shelf_var.get())
            messagebox.showinfo("Başarılı", "Kitap eklendi.")
            self.show_home()

        def import_file():
            file_name = filedialog.askopenfilename(
                title="Kitapları içe aktar",
                filetypes=[("Excel/CSV", "*.xlsx *.xlsm *.csv"), ("Tüm Dosyalar", "*.*")],
            )
            if not file_name:
                return
            try:
                mapping = self._open_import_mapping_dialog(Path(file_name))
                if not mapping:
                    return
                count = 0
                def _do_import():
                    nonlocal count
                    count = self.data.import_books(Path(file_name), mapping)
                self._run_with_progress("İçe Aktarma", _do_import)
            except Exception as err:
                messagebox.showerror("İçe aktarma hatası", str(err))
                return
            messagebox.showinfo("Başarılı", f"{count} kitap içe aktarıldı.")
            self.show_home()

        actions = ttk.Frame(wrapper)
        actions.pack(anchor="w", pady=8)
        ttk.Button(actions, text="Kaydet", command=save).grid(row=0, column=0, padx=(0, 8))
        ttk.Button(actions, text="Excel/CSV İçe Aktar", command=import_file).grid(row=0, column=1)

    def show_manage_book(self):
        self._clear_main()
        wrapper = ttk.Frame(self.main_frame, padding=20)
        wrapper.grid(row=0, column=0, sticky="nsew")
        self._standard_form_title(wrapper, "Kitap Düzenle / Sil", "ID ile getirip bilgileri güncelleyebilir veya silebilirsin")

        code_var, title_var, author_var, shelf_var = tk.StringVar(), tk.StringVar(), tk.StringVar(), tk.StringVar()
        form = ttk.Frame(wrapper)
        form.pack(anchor="w", pady=8)

        for i, (label, var) in enumerate([("QR / Kitap ID", code_var), ("Kitap Adı", title_var), ("Yazar", author_var), ("Raf", shelf_var)]):
            ttk.Label(form, text=f"{label}:").grid(row=i, column=0, sticky="w", pady=5)
            ttk.Entry(form, textvariable=var, width=42).grid(row=i, column=1, sticky="w", padx=8, pady=5)

        def load_by_id():
            try:
                row = self.data._find_row(self.data.read(), code_var.get())
                if not row:
                    raise ValueError("Kitap bulunamadı.")
            except Exception as err:
                messagebox.showerror("Hata", str(err))
                return
            title_var.set(row["KitapAdi"])
            author_var.set(row["Yazar"])
            shelf_var.set(row["Raf"])

        def update_book():
            try:
                self._run_with_progress("Güncelleme", lambda: self.data.update_book(code_var.get(), title_var.get(), author_var.get(), shelf_var.get()))
            except Exception as err:
                messagebox.showerror("Hata", str(err))
                return
            messagebox.showinfo("Başarılı", "Kitap güncellendi.")
            self.show_home()

        def delete_book():
            if not messagebox.askyesno("Onay", "Bu kitabı silmek istediğinize emin misiniz?"):
                return
            try:
                self._run_with_progress("Silme", lambda: self.data.delete_book(code_var.get()))
            except Exception as err:
                messagebox.showerror("Hata", str(err))
                return
            messagebox.showinfo("Başarılı", "Kitap silindi.")
            self.show_home()

        action = ttk.Frame(wrapper)
        action.pack(anchor="w", pady=6)
        ttk.Button(action, text="QR/ID ile Getir", command=load_by_id).grid(row=0, column=0, padx=(0, 6))
        ttk.Button(action, text="Güncelle", command=update_book).grid(row=0, column=1, padx=6)
        ttk.Button(action, text="Sil", command=delete_book).grid(row=0, column=2, padx=6)

    def show_students_panel(self):
        self._clear_main()
        wrapper = ttk.Frame(self.main_frame, padding=20)
        wrapper.grid(row=0, column=0, sticky="nsew")
        self._standard_form_title(wrapper, "Öğrenci Kontrol Paneli", "Öğrenci numarası ve isim ile kaydet/güncelle/sil")

        top = ttk.Frame(wrapper)
        top.pack(fill='x', pady=8)
        search_var = tk.StringVar()
        ttk.Label(top, text='Ara (numara/isim):').pack(side='left')
        ttk.Entry(top, textvariable=search_var, width=34).pack(side='left', padx=8)

        body = ttk.Frame(wrapper)
        body.pack(fill='both', expand=True)
        body.columnconfigure(1, weight=1)
        body.rowconfigure(0, weight=1)

        form = ttk.LabelFrame(body, text='Öğrenci Bilgisi', padding=12)
        form.grid(row=0, column=0, sticky='nsw', padx=(0, 12))

        no_var = tk.StringVar()
        name_var = tk.StringVar()

        ttk.Label(form, text='Öğrenci No:').grid(row=0, column=0, sticky='w', pady=6)
        ttk.Entry(form, textvariable=no_var, width=20).grid(row=0, column=1, pady=6)

        ttk.Label(form, text='Ad Soyad:').grid(row=1, column=0, sticky='w', pady=6)
        ttk.Entry(form, textvariable=name_var, width=26).grid(row=1, column=1, pady=6)

        tree = ttk.Treeview(body, columns=STUDENT_COLUMNS, show='headings', height=18)
        tree.heading('OgrenciNo', text='Öğrenci No')
        tree.heading('AdSoyad', text='Ad Soyad')
        tree.column('OgrenciNo', width=130, anchor='center')
        tree.column('AdSoyad', width=320, anchor='w')
        tree.grid(row=0, column=1, sticky='nsew')

        def load_students(*_):
            tree.delete(*tree.get_children())
            q = search_var.get().strip().lower()
            for r in self.students.read():
                text = f"{r['OgrenciNo']} {r['AdSoyad']}".lower()
                if q and q not in text:
                    continue
                tree.insert('', 'end', values=(r['OgrenciNo'], r['AdSoyad']))

        def clear_form():
            no_var.set('')
            name_var.set('')

        def add_student():
            if not no_var.get().strip() or not name_var.get().strip():
                messagebox.showwarning('Eksik bilgi', 'Öğrenci numarası ve adı boş olamaz.')
                return
            try:
                new_no = None
                def _action():
                    nonlocal new_no
                    new_no = self.students.add_student(no_var.get(), name_var.get())
                self._run_with_progress('Öğrenci Kaydı', _action)
            except Exception as err:
                messagebox.showerror('Hata', str(err))
                return
            messagebox.showinfo('Başarılı', f'Öğrenci kaydedildi. No: {new_no}')
            clear_form()
            load_students()

        def update_student():
            if not no_var.get().strip():
                messagebox.showwarning('Uyarı', 'Güncellemek için öğrenci no girin veya listeden seçin.')
                return
            if not name_var.get().strip():
                messagebox.showwarning('Eksik bilgi', 'Öğrenci adı boş olamaz.')
                return
            try:
                self._run_with_progress('Öğrenci Güncelleme', lambda: self.students.update_student(no_var.get(), name_var.get()))
            except Exception as err:
                messagebox.showerror('Hata', str(err))
                return
            messagebox.showinfo('Başarılı', 'Öğrenci güncellendi.')
            load_students()

        def delete_student():
            if not no_var.get().strip():
                messagebox.showwarning('Uyarı', 'Silmek için öğrenci no girin veya listeden seçin.')
                return
            if not messagebox.askyesno('Onay', 'Seçili öğrenciyi silmek istiyor musunuz?'):
                return
            try:
                self._run_with_progress('Öğrenci Silme', lambda: self.students.delete_student(no_var.get()))
            except Exception as err:
                messagebox.showerror('Hata', str(err))
                return
            messagebox.showinfo('Başarılı', 'Öğrenci silindi.')
            clear_form()
            load_students()

        btns = ttk.Frame(form)
        btns.grid(row=2, column=0, columnspan=2, pady=(10, 0), sticky='w')
        ttk.Button(btns, text='Kaydet', command=add_student).grid(row=0, column=0, padx=(0, 6))
        ttk.Button(btns, text='Güncelle', command=update_student).grid(row=0, column=1, padx=6)
        ttk.Button(btns, text='Sil', command=delete_student).grid(row=0, column=2, padx=6)
        ttk.Button(btns, text='Temizle', command=clear_form).grid(row=0, column=3, padx=6)

        def on_select(_event=None):
            item = tree.focus()
            if not item:
                return
            vals = tree.item(item, 'values')
            if len(vals) >= 2:
                no_var.set(vals[0])
                name_var.set(vals[1])

        tree.bind('<<TreeviewSelect>>', on_select)
        search_var.trace_add('write', load_students)
        load_students()

    def show_lend(self):
        self._clear_main()
        wrapper = ttk.Frame(self.main_frame, padding=20)
        wrapper.grid(row=0, column=0, sticky="nsew")
        self._standard_form_title(wrapper, "Kitap Ödünç Ver", "Teslim tarihi otomatik hesaplanır")

        code_var, borrower_var, days_var = tk.StringVar(), tk.StringVar(), tk.StringVar(value="14")
        form = ttk.Frame(wrapper)
        form.pack(anchor="w", pady=10)

        for i, (label, var) in enumerate([("QR / Kitap ID", code_var), ("Ödünç Alan", borrower_var), ("Kaç Gün", days_var)]):
            ttk.Label(form, text=f"{label}:").grid(row=i, column=0, sticky="w", pady=6)
            ttk.Entry(form, textvariable=var, width=38).grid(row=i, column=1, padx=8, pady=6)

        def lend():
            try:
                if not borrower_var.get().strip():
                    raise ValueError("Ödünç alan boş olamaz.")
                days = int(days_var.get())
                if days < 1:
                    raise ValueError("Gün sayısı en az 1 olmalı.")
                self._run_with_progress("Ödünç Verme", lambda: self.data.lend_book(code_var.get(), borrower_var.get(), days))
            except Exception as err:
                messagebox.showerror("Hata", str(err))
                return
            messagebox.showinfo("Başarılı", "Kitap ödünç verildi.")
            self.show_home()

        ttk.Button(wrapper, text="Ödünç Ver", command=lend).pack(anchor="w", pady=8)

    def show_return(self):
        self._clear_main()
        wrapper = ttk.Frame(self.main_frame, padding=20)
        wrapper.grid(row=0, column=0, sticky="nsew")
        self._standard_form_title(wrapper, "Kitap İade", "QR ya da ID girerek hızlı iade işlemi yap")

        code_var = tk.StringVar()
        row = ttk.Frame(wrapper)
        row.pack(anchor="w", pady=12)
        ttk.Label(row, text="QR / Kitap ID:").pack(side="left")
        ttk.Entry(row, textvariable=code_var, width=24).pack(side="left", padx=8)

        def process_return():
            try:
                self._run_with_progress("İade Alma", lambda: self.data.return_book(code_var.get()))
            except Exception as err:
                messagebox.showerror("Hata", str(err))
                return
            messagebox.showinfo("Başarılı", "Kitap iade alındı.")
            self.show_home()

        ttk.Button(wrapper, text="İade Al", command=process_return).pack(anchor="w", pady=6)

    def show_extend_due(self):
        self._clear_main()
        wrapper = ttk.Frame(self.main_frame, padding=20)
        wrapper.grid(row=0, column=0, sticky="nsew")
        self._standard_form_title(wrapper, "Teslim Süresi Uzat", "Ödünçteki kitaplar için teslim tarihine gün ekler")

        code_var, days_var = tk.StringVar(), tk.StringVar(value="7")
        form = ttk.Frame(wrapper)
        form.pack(anchor="w", pady=10)
        ttk.Label(form, text="QR / Kitap ID:").grid(row=0, column=0, sticky="w", pady=6)
        ttk.Entry(form, textvariable=code_var, width=24).grid(row=0, column=1, padx=8, pady=6)
        ttk.Label(form, text="Eklenecek Gün:").grid(row=1, column=0, sticky="w", pady=6)
        ttk.Entry(form, textvariable=days_var, width=24).grid(row=1, column=1, padx=8, pady=6)

        def extend_due():
            try:
                days = int(days_var.get())
                if days < 1:
                    raise ValueError("Eklenecek gün en az 1 olmalı.")
                self._run_with_progress("Süre Uzatma", lambda: self.data.extend_due_date(code_var.get(), days))
            except Exception as err:
                messagebox.showerror("Hata", str(err))
                return
            messagebox.showinfo("Başarılı", "Teslim tarihi güncellendi.")
            self.show_home()

        ttk.Button(wrapper, text="Süre Uzat", command=extend_due).pack(anchor="w", pady=6)

    def show_overdues(self):
        self._clear_main()
        wrapper = ttk.Frame(self.main_frame, padding=20)
        wrapper.grid(row=0, column=0, sticky="nsew")
        self._standard_form_title(wrapper, "Geciken Kitaplar", "Teslim tarihi geçen kayıtlar")

        cols = ["KitapID", "KitapAdi", "OduncAlan", "TeslimTarihi", "GecikmeGun"]
        tree = ttk.Treeview(wrapper, columns=cols, show="headings", height=18)
        for col in cols:
            tree.heading(col, text=col)
            tree.column(col, width=170, anchor="center")
        tree.pack(fill="both", expand=True, pady=(8, 0))

        for row, delay in self.data.overdue_books():
            tree.insert("", "end", values=(row["KitapID"], row["KitapAdi"], row["OduncAlan"], row["TeslimTarihi"], delay))

    def make_backup(self):
        try:
            target = self.data.create_backup()
        except Exception as err:
            messagebox.showerror("Hata", str(err))
            return
        messagebox.showinfo("Yedek alındı", f"Yedek dosyası oluşturuldu:\n{target}")


    def _handle_db_clear_shortcut(self, _event=None):
        self.clear_database_action()

    def clear_database_action(self):
        if not messagebox.askyesno("Veritabanını Temizle", "Tüm kitap kayıtları silinecek. Devam edilsin mi?"):
            return
        try:
            self._run_with_progress("Veritabanı Temizleniyor", self.data.clear_database)
        except Exception as err:
            messagebox.showerror("Hata", str(err))
            return
        messagebox.showinfo("Tamamlandı", "Veritabanı temizlendi.")
        self.show_home()

    def exit_app(self):
        if messagebox.askyesno("Çıkış", "Uygulamadan çıkmak istiyor musunuz?"):
            self.destroy()


if __name__ == "__main__":
    app = LibraryApp()
    app.mainloop()
